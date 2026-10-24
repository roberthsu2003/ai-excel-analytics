"""產生第一部「用 AI 生成 Excel」各課的觀察用 xlsx。

每個 build_XX() 對應一課，完全依照該課資料夾內的規範書產生檔案。
執行：uv run python scripts/build_part1_workbooks.py
"""

from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

ROOT = Path(__file__).resolve().parent.parent / "01_用AI生成Excel"

FONT_NAME = "Microsoft JhengHei"
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(name=FONT_NAME, bold=True, color="FFFFFF")
LABEL_FILL = PatternFill("solid", fgColor="DDEBF7")
BODY_FONT = Font(name=FONT_NAME)
NOTE_FONT = Font(name=FONT_NAME, italic=True, color="595959")

AMOUNT = "#,##0"
DATE = "yyyy-mm-dd"


def write_header(ws, row, headers):
    for col, text in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=col, value=text)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")


def write_rows(ws, start_row, rows, formats=None):
    """寫入資料列；formats 是 {欄號: 數字格式}。"""
    formats = formats or {}
    for r, values in enumerate(rows, start=start_row):
        for c, value in enumerate(values, start=1):
            cell = ws.cell(row=r, column=c, value=value)
            cell.font = BODY_FONT
            if c in formats:
                cell.number_format = formats[c]


def set_widths(ws, widths):
    for col, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(col)].width = width


def label(ws, ref, text):
    cell = ws[ref]
    cell.value = text
    cell.font = Font(name=FONT_NAME, bold=True)
    cell.fill = LABEL_FILL


def add_table(ws, ref, name):
    table = Table(displayName=name, ref=ref)
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(table)


def lesson_path(folder, filename):
    return ROOT / folder / filename


# ---------------------------------------------------------------- 01
def build_01():
    wb = Workbook()
    ws = wb.active
    ws.title = "收支紀錄"
    headers = ["日期", "收支類別", "項目說明", "金額", "經手人", "收據編號"]
    rows = [
        (date(2026, 2, 23), "收入", "社員社費（25 人 × 300 元）", 7500, "林雨萱", "0001"),
        (date(2026, 2, 25), "支出", "社課講師鐘點費", 1600, "陳柏宇", "0002"),
        (date(2026, 3, 2), "支出", "相機電池與記憶卡", 2380, "陳柏宇", "0003"),
        (date(2026, 3, 9), "收入", "學務處社團補助", 5000, "林雨萱", "0004"),
        (date(2026, 3, 16), "支出", "外拍交通費（客運）", 1250, "王子晴", "0005"),
        (date(2026, 3, 23), "支出", "攝影展海報輸出", 860, "王子晴", "0006"),
        (date(2026, 3, 30), "收入", "攝影展明信片義賣", 3420, "林雨萱", "0007"),
        (date(2026, 4, 6), "支出", "期中聚餐補貼", 1800, "陳柏宇", "0008"),
        (date(2026, 4, 13), "支出", "器材維修（腳架）", 650, "王子晴", "0009"),
        (date(2026, 4, 20), "收入", "校友贊助", 2000, "林雨萱", "0010"),
    ]
    write_header(ws, 1, headers)
    write_rows(ws, 2, rows, {1: DATE, 4: AMOUNT, 6: "@"})
    set_widths(ws, [13, 11, 30, 11, 10, 11])
    ws.freeze_panes = "A2"

    demo = wb.create_sheet("型態對照")
    write_header(demo, 1, ["比較項目", "存成文字（錯誤）", "存成正確型態", "說明"])
    demo_rows = [
        ("金額 1", "1500", 1500, "文字的 1500 靠左，數值的 1500 靠右"),
        ("金額 2", "2,000元", 2000, "數值不打「元」，千分位交給格式 #,##0"),
        ("金額 3", "NT$3500", 3500, "貨幣符號也不要打進儲存格"),
        ("加總", "=SUM(B2:B4)", "=SUM(C2:C4)", "文字欄加總 = 0；數值欄加總 = 7,000"),
        ("日期", "115.03.02", date(2026, 3, 2), "民國年加句點是文字；右邊是真正的日期"),
        ("取出月份", "=MONTH(B6)", "=MONTH(C6)", "文字日期 → #VALUE!；真日期 → 3"),
        ("收據編號", 12, "0012", "編號要保留前導 0，所以存成「文字」"),
    ]
    write_rows(demo, 2, demo_rows)
    for ref in ("C2", "C3", "C4", "C5"):
        demo[ref].number_format = AMOUNT
    demo["C6"].number_format = DATE
    demo["C8"].number_format = "@"
    set_widths(demo, [12, 18, 16, 42])
    demo.freeze_panes = "A2"

    wb.save(lesson_path("01_儲存格與資料型態", "社團收支紀錄.xlsx"))


# ---------------------------------------------------------------- 02
QUOTE_ITEMS = [
    ("P-001", "原子筆（藍，12 支裝）", 120, 10),
    ("P-002", "A4 影印紙（500 張）", 135, 20),
    ("P-003", "釘書機", 180, 5),
    ("P-004", "便利貼（76×76 mm）", 45, 30),
    ("P-005", "白板筆（4 色組）", 160, 8),
    ("P-006", "A4 資料夾（10 入）", 95, 15),
    ("P-007", "計算機", 450, 3),
    ("P-008", "剪刀", 70, 6),
]


def _quote_sheet(ws, tax_formula, with_totals=True):
    """報價單版面：上方參數區（A1:B4），第 6 列表頭，第 7 列起明細。"""
    label(ws, "A1", "營業稅率")
    ws["B1"] = 0.05
    ws["B1"].number_format = "0%"
    if with_totals:
        _quote_totals(ws)
    _quote_details(ws, tax_formula)


def _quote_totals(ws):
    label(ws, "A2", "未稅合計")
    ws["B2"] = "=SUM(E7:E14)"
    label(ws, "A3", "稅額合計")
    ws["B3"] = "=SUM(F7:F14)"
    label(ws, "A4", "含稅總額")
    ws["B4"] = "=B2+B3"
    for ref in ("B2", "B3", "B4"):
        ws[ref].number_format = AMOUNT
        ws[ref].font = Font(name=FONT_NAME, bold=True)


def _quote_details(ws, tax_formula):
    write_header(ws, 6, ["品項編號", "品名", "單價", "數量", "小計", "稅額", "含稅金額"])
    rows = []
    for i, (code, name, price, qty) in enumerate(QUOTE_ITEMS):
        r = 7 + i
        rows.append((code, name, price, qty, f"=C{r}*D{r}", tax_formula(r), f"=E{r}+F{r}"))
    write_rows(ws, 7, rows, {3: AMOUNT, 4: AMOUNT, 5: AMOUNT, 6: AMOUNT, 7: AMOUNT})
    set_widths(ws, [12, 24, 10, 8, 11, 10, 12])
    ws.freeze_panes = "A7"


def build_02():
    wb = Workbook()
    ws = wb.active
    ws.title = "報價單"
    # 正確：稅率格用絕對參照 $B$1，向下複製時不會跑掉
    _quote_sheet(ws, lambda r: f"=ROUND(E{r}*$B$1,0)")

    wrong = wb.create_sheet("錯誤示範_相對參照")
    # 錯誤：第 7 列寫 =E7*B1 後向下填滿，第 8 列變成 =E8*B2（空白）、第 12 列起乘到文字
    # 不放合計，避免錯位的參照指到合計格而形成循環參照
    _quote_sheet(wrong, lambda r: f"=ROUND(E{r}*B{r - 6},0)", with_totals=False)
    wrong["D1"] = "⚠️ F 欄稅額用了相對參照 B1，向下填滿後變成 B2、B3…，數字全錯"
    wrong["D1"].font = Font(name=FONT_NAME, bold=True, color="C00000")

    wb.save(lesson_path("02_公式與參照", "商品報價單.xlsx"))


# ---------------------------------------------------------------- 03
STUDENTS = [
    ("01", "陳怡君", 82, 76, 91), ("02", "林家豪", 68, 72, 55), ("03", "黃詩涵", 90, 88, 84),
    ("04", "張志明", 74, None, 63), ("05", "李宜蓁", 88, 93, 79), ("06", "王俊傑", 59, 64, 71),
    ("07", "吳佩珊", 95, 89, 92), ("08", "劉冠廷", 71, 58, 66), ("09", "蔡雅婷", 84, 80, 88),
    ("10", "楊承恩", 66, 70, None), ("11", "許書瑜", 77, 85, 73), ("12", "鄭宇翔", 62, 55, 48),
    ("13", "謝欣妤", 91, 94, 87), ("14", "郭品睿", 80, 67, 82), ("15", "洪子晴", 73, 79, 69),
    ("16", "曾柏翰", 57, 61, 52), ("17", "邱于庭", 86, 82, 90), ("18", "廖哲瑋", 69, 74, 77),
    ("19", "賴思妤", 93, 90, 95), ("20", "周彥廷", 78, 71, 64),
]


def build_03():
    wb = Workbook()
    ws = wb.active
    ws.title = "成績"
    write_header(ws, 1, ["座號", "姓名", "國文", "英文", "數學", "總分", "平均", "已考科數"])
    rows = []
    for i, (no, name, chi, eng, math) in enumerate(STUDENTS):
        r = 2 + i
        rows.append((no, name, chi, eng, math,
                     f"=SUM(C{r}:E{r})", f"=ROUND(AVERAGE(C{r}:E{r}),1)", f"=COUNT(C{r}:E{r})"))
    write_rows(ws, 2, rows, {1: "@", 7: "0.0"})
    set_widths(ws, [7, 10, 8, 8, 8, 8, 8, 10])
    ws.freeze_panes = "C2"

    stats = wb.create_sheet("各科統計")
    write_header(stats, 1, ["統計項目", "國文", "英文", "數學", "使用的函數"])
    last = 1 + len(STUDENTS)
    items = [
        ("平均分數", "=ROUND(AVERAGE('成績'!{c}2:{c}%d),1)" % last, "AVERAGE＋ROUND"),
        ("最高分", "=MAX('成績'!{c}2:{c}%d)" % last, "MAX"),
        ("最低分", "=MIN('成績'!{c}2:{c}%d)" % last, "MIN"),
        ("應考人數", "=COUNT('成績'!{c}2:{c}%d)" % last, "COUNT（只數數字）"),
        ("全班人數", "=COUNTA('成績'!B2:B%d)" % last, "COUNTA（數非空白）"),
    ]
    for r, (name, pattern, func) in enumerate(items, start=2):
        stats.cell(row=r, column=1, value=name).font = Font(name=FONT_NAME, bold=True)
        for c, col in enumerate("CDE", start=2):
            cell = stats.cell(row=r, column=c, value=pattern.format(c=col))
            cell.font = BODY_FONT
            if r == 2:
                cell.number_format = "0.0"
        stats.cell(row=r, column=5, value=func).font = NOTE_FONT
    set_widths(stats, [12, 9, 9, 9, 22])

    wb.save(lesson_path("03_常用函數", "班級成績表.xlsx"))


# ---------------------------------------------------------------- 04
EXPENSES = [
    ("EXP-2026-001", date(2026, 3, 2), "研發部", "林志遠", "交通差旅", "高鐵來回車票（台北－新竹客戶端技術對接）", 580, 2, "AB-10293847", "附票根"),
    ("EXP-2026-002", date(2026, 3, 3), "行銷部", "張雅晴", "業務推廣", "春季新品發表會場地租借", 12000, 1, "AC-99887766", "附合約"),
    ("EXP-2026-003", date(2026, 3, 5), "維運部", "陳建宏", "硬體耗材", "機房網路線與轉接頭", 450, 6, "AD-33445566", "已入庫"),
    ("EXP-2026-004", date(2026, 3, 8), "研發部", "王思婷", "軟體訂閱", "雲端協作平台團隊月租", 1800, 1, "AE-55667788", ""),
    ("EXP-2026-005", date(2026, 3, 10), "業務部", "黃俊傑", "交際應酬", "策略夥伴商務餐敘", 4800, 1, "AF-77889900", "附簽到單"),
    ("EXP-2026-006", date(2026, 3, 12), "管理部", "劉美琪", "辦公行政", "季度文具與事務耗材", 350, 12, "AG-11223344", ""),
    ("EXP-2026-007", date(2026, 3, 15), "行銷部", "張雅晴", "廣告投放", "數位影音廣告春季預算", 25000, 1, "AH-66778899", "附結案報告"),
    ("EXP-2026-008", date(2026, 3, 18), "研發部", "林志遠", "專業培訓", "雲端架構證照考試報名費", 6500, 2, "AI-44332211", ""),
    ("EXP-2026-009", date(2026, 3, 22), "業務部", "黃俊傑", "交通差旅", "業務拜訪油資與過路費", 1450, 1, "AJ-88990011", "附里程表"),
    ("EXP-2026-010", date(2026, 3, 25), "管理部", "劉美琪", "會議雜費", "季會茶點", 280, 15, "AK-22334455", ""),
]


def build_04():
    wb = Workbook()
    ws = wb.active
    ws.title = "報銷明細"
    headers = ["報銷序號", "報銷日期", "申請部門", "申請人", "費用類別", "支出項目說明",
               "單價", "數量", "小計金額", "發票號碼", "審核備註"]
    write_header(ws, 1, headers)
    rows = []
    for i, (no, d, dept, who, cat, desc, price, qty, inv, note) in enumerate(EXPENSES):
        r = 2 + i
        rows.append((no, d, dept, who, cat, desc, price, qty, f"=G{r}*H{r}", inv, note))
    write_rows(ws, 2, rows, {2: DATE, 7: AMOUNT, 8: AMOUNT, 9: AMOUNT})
    for r in range(2, 2 + len(rows)):
        ws.cell(row=r, column=6).alignment = Alignment(wrap_text=True, vertical="top")
    last = 1 + len(rows)
    add_table(ws, f"A1:K{last}", "T_Expense")
    set_widths(ws, [15, 12, 10, 9, 11, 32, 9, 7, 11, 14, 12])
    ws.freeze_panes = "A2"

    summary = wb.create_sheet("報銷彙總")
    write_header(summary, 1, ["彙總項目", "數值", "公式說明"])
    items = [
        ("報銷筆數", "=COUNTA(T_Expense[報銷序號])", "數表格中有幾個報銷序號", "0"),
        ("報銷總計金額", "=SUM(T_Expense[小計金額])", "加總整欄小計，新增列會自動算進來", AMOUNT),
        ("最大單筆金額", "=MAX(T_Expense[小計金額])", "找出最大的一筆", AMOUNT),
    ]
    for r, (name, formula, note, fmt) in enumerate(items, start=2):
        summary.cell(row=r, column=1, value=name).font = Font(name=FONT_NAME, bold=True)
        cell = summary.cell(row=r, column=2, value=formula)
        cell.number_format = fmt
        cell.font = BODY_FONT
        summary.cell(row=r, column=3, value=note).font = NOTE_FONT
    set_widths(summary, [16, 14, 36])

    wb.save(lesson_path("04_一列一筆與結構化表格", "部門費用報銷表.xlsx"))


# ---------------------------------------------------------------- 05
SALES_REPS = [
    ("E001", "陳怡君", "業務一部", 500000, 612000, "是"),
    ("E002", "林家豪", "業務一部", 450000, 452000, "否"),
    ("E003", "黃詩涵", "業務一部", 480000, 384000, "是"),
    ("E004", "張志明", "業務二部", 520000, 624000, "否"),
    ("E005", "李宜蓁", "業務二部", 400000, 300000, "是"),
    ("E006", "王俊傑", "業務二部", 500000, 500000, "是"),
    ("E007", "吳佩珊", "業務三部", 600000, 690000, "是"),
    ("E008", "劉冠廷", "業務三部", 450000, 396000, "否"),
    ("E009", "蔡雅婷", "業務三部", 380000, 471200, "是"),
    ("E010", "楊承恩", "業務一部", 420000, 344400, "否"),
    ("E011", "許書瑜", "業務二部", 550000, 412500, "是"),
    ("E012", "鄭宇翔", "業務三部", 500000, 590000, "否"),
]
BONUS_RULES = [("A", 1.2, "達成率 ≥ 120%"), ("B", 1.0, "達成率 ≥ 100%"), ("C", 0.8, "達成率 ≥ 80%")]


def build_05():
    wb = Workbook()
    ws = wb.active
    ws.title = "業績"
    write_header(ws, 1, ["員工編號", "姓名", "部門", "業績目標", "實際業績", "全勤",
                         "達成率", "是否達標", "獎金等級", "全勤加碼"])
    rule = "'獎金規則'!"
    rows = []
    for i, rep in enumerate(SALES_REPS):
        r = 2 + i
        rows.append(rep + (
            f"=E{r}/D{r}",
            f'=IF(G{r}>={rule}$B$3,"達標","未達標")',
            # IFS 是 Excel 2019 起的新函數，檔案內必須寫成 _xlfn.IFS
            f'=_xlfn.IFS(G{r}>={rule}$B$2,"A",G{r}>={rule}$B$3,"B",G{r}>={rule}$B$4,"C",TRUE,"無")',
            f'=IF(AND(H{r}="達標",F{r}="是"),"是","否")',
        ))
    write_rows(ws, 2, rows, {4: AMOUNT, 5: AMOUNT, 7: "0.0%"})
    add_table(ws, f"A1:J{1 + len(rows)}", "T_Bonus")
    set_widths(ws, [10, 9, 10, 11, 11, 7, 9, 10, 10, 10])
    ws.freeze_panes = "C2"

    rules = wb.create_sheet("獎金規則")
    write_header(rules, 1, ["獎金等級", "達成率下限", "說明"])
    write_rows(rules, 2, BONUS_RULES + [("無", None, "達成率未滿 80%")], {2: "0%"})
    rules["A7"] = "全勤加碼：「是否達標」為達標，而且「全勤」為是，兩個條件都要成立（AND）。"
    rules["A7"].font = NOTE_FONT
    set_widths(rules, [10, 12, 20])

    wb.save(lesson_path("05_邏輯判斷", "業績獎金判定.xlsx"))


# ---------------------------------------------------------------- 06
LEAVE_OPTIONS = {
    "部門": ["研發部", "行銷部", "維運部", "業務部", "管理部"],
    "假別": ["特休", "事假", "病假", "公假", "婚假", "喪假"],
    "審核狀態": ["待審核", "核准", "駁回"],
}
LEAVES = [
    ("L-001", date(2026, 3, 2), "林志遠", "研發部", "特休", date(2026, 3, 9), date(2026, 3, 10), 2, "核准", ""),
    ("L-002", date(2026, 3, 3), "張雅晴", "行銷部", "病假", date(2026, 3, 3), date(2026, 3, 3), 1, "核准", ""),
    ("L-003", date(2026, 3, 5), "陳建宏", "維運部", "事假", date(2026, 3, 12), date(2026, 3, 12), 0.5, "待審核", "下午半天"),
    ("L-004", date(2026, 3, 6), "王思婷", "研發部", "特休", date(2026, 3, 23), date(2026, 3, 27), 5, "待審核", ""),
    ("L-005", date(2026, 3, 10), "黃俊傑", "業務部", "喪假", date(2026, 3, 11), date(2026, 3, 13), 3, "核准", ""),
    ("L-006", date(2026, 3, 12), "劉美琪", "管理部", "公假", date(2026, 3, 18), date(2026, 3, 18), 1, "核准", "教育訓練"),
    ("L-007", date(2026, 3, 15), "張雅晴", "行銷部", "特休", date(2026, 4, 1), date(2026, 4, 3), 3, "駁回", "與發表會衝突，請改期"),
    ("L-008", date(2026, 3, 18), "陳建宏", "維運部", "病假", date(2026, 3, 17), date(2026, 3, 18), 2, "核准", "事後補請"),
    ("L-009", date(2026, 3, 20), "林志遠", "研發部", "事假", date(2026, 3, 30), date(2026, 3, 30), 1, "待審核", ""),
    ("L-010", date(2026, 3, 25), "黃俊傑", "業務部", "婚假", date(2026, 4, 20), date(2026, 4, 24), 5, "待審核", ""),
]
VALIDATION_ROWS = 200  # 驗證套用到第 200 列，之後新增的資料也受保護


def _validation(kind, sqref, prompt, error, **kwargs):
    from openpyxl.worksheet.datavalidation import DataValidation

    dv = DataValidation(type=kind, allow_blank=True, showInputMessage=True, showErrorMessage=True,
                        promptTitle="填寫說明", prompt=prompt, errorTitle="輸入錯誤", error=error,
                        errorStyle="stop", **kwargs)
    dv.add(sqref)
    return dv


def build_06():
    wb = Workbook()
    ws = wb.active
    ws.title = "請假紀錄"
    write_header(ws, 1, ["申請編號", "申請日期", "員工姓名", "部門", "假別", "開始日期",
                         "結束日期", "請假天數", "審核狀態", "備註"])
    write_rows(ws, 2, LEAVES, {2: DATE, 6: DATE, 7: DATE, 8: "0.0"})
    add_table(ws, f"A1:J{1 + len(LEAVES)}", "T_Leave")
    set_widths(ws, [10, 12, 10, 9, 8, 12, 12, 10, 10, 22])
    ws.freeze_panes = "A2"

    opts = wb.create_sheet("選項清單")
    write_header(opts, 1, list(LEAVE_OPTIONS))
    for c, values in enumerate(LEAVE_OPTIONS.values(), start=1):
        for r, v in enumerate(values, start=2):
            opts.cell(row=r, column=c, value=v).font = BODY_FONT
    set_widths(opts, [10, 10, 10])

    n = VALIDATION_ROWS
    ranges = {k: f"'選項清單'!${get_column_letter(i)}$2:${get_column_letter(i)}${1 + len(v)}"
              for i, (k, v) in enumerate(LEAVE_OPTIONS.items(), start=1)}
    for dv in [
        _validation("list", f"D2:D{n}", "請從下拉選單選擇部門", "部門只能從選單中選擇", formula1=ranges["部門"]),
        _validation("list", f"E2:E{n}", "請從下拉選單選擇假別", "假別只能從選單中選擇", formula1=ranges["假別"]),
        _validation("list", f"I2:I{n}", "待審核／核准／駁回", "審核狀態只能從選單中選擇", formula1=ranges["審核狀態"]),
        _validation("date", f"F2:F{n}", "請輸入 2026 年的日期", "開始日期必須在 2026-01-01 ～ 2026-12-31 之間",
                    operator="between", formula1="DATE(2026,1,1)", formula2="DATE(2026,12,31)"),
        _validation("custom", f"G2:G{n}", "結束日期不可早於開始日期", "結束日期不可早於開始日期", formula1="G2>=F2"),
        _validation("decimal", f"H2:H{n}", "以 0.5 天為單位，最多 14 天", "請假天數必須介於 0.5 ～ 14 天",
                    operator="between", formula1="0.5", formula2="14"),
    ]:
        ws.add_data_validation(dv)

    wb.save(lesson_path("06_資料驗證", "請假申請表.xlsx"))


# ---------------------------------------------------------------- 07
BASE_DATE = date(2026, 4, 15)
TASKS = [
    ("T-01", "需求訪談", "王思婷", date(2026, 2, 28), "已完成", 1.0, 30000, 28000),
    ("T-02", "系統架構設計", "林志遠", date(2026, 3, 15), "已完成", 1.0, 80000, 86000),
    ("T-03", "UI 設計稿", "張雅晴", date(2026, 3, 31), "進行中", 0.9, 60000, 45000),
    ("T-04", "前端開發", "陳建宏", date(2026, 5, 15), "進行中", 0.6, 150000, 120000),
    ("T-05", "後端 API 開發", "林志遠", date(2026, 5, 31), "進行中", 0.7, 180000, 192000),
    ("T-06", "資料庫建置", "陳建宏", date(2026, 4, 10), "進行中", 0.8, 70000, 52000),
    ("T-07", "測試計畫", "劉美琪", date(2026, 4, 30), "進行中", 0.3, 25000, 8000),
    ("T-08", "使用者測試", "王思婷", date(2026, 6, 15), "未開始", 0.0, 40000, 0),
    ("T-09", "教育訓練教材", "劉美琪", date(2026, 6, 10), "進行中", 0.2, 35000, 12000),
    ("T-10", "上線部署", "陳建宏", date(2026, 6, 30), "未開始", 0.0, 50000, 0),
    ("T-11", "行銷素材", "張雅晴", date(2026, 4, 5), "已完成", 1.0, 45000, 47500),
    ("T-12", "驗收報告", "黃俊傑", date(2026, 7, 15), "未開始", 0.0, 20000, 0),
]


def build_07():
    from openpyxl.formatting.rule import CellIsRule, DataBarRule, FormulaRule
    from openpyxl.workbook.defined_name import DefinedName

    wb = Workbook()
    ws = wb.active
    ws.title = "專案進度"
    write_header(ws, 1, ["任務編號", "任務名稱", "負責人", "截止日", "狀態", "完成進度",
                         "預算", "實際支出", "剩餘預算", "預算使用率"])
    rows = []
    for i, task in enumerate(TASKS):
        r = 2 + i
        rows.append(task + (f"=G{r}-H{r}", f"=IF(G{r}=0,0,H{r}/G{r})"))
    write_rows(ws, 2, rows, {4: DATE, 6: "0%", 7: AMOUNT, 8: AMOUNT, 9: AMOUNT, 10: "0%"})
    last = 1 + len(rows)
    add_table(ws, f"A1:J{last}", "T_Project")
    set_widths(ws, [10, 16, 9, 12, 9, 10, 11, 11, 11, 11])
    ws.freeze_panes = "C2"

    setting = wb.create_sheet("設定")
    label(setting, "A1", "基準日")
    setting["B1"] = BASE_DATE
    setting["B1"].number_format = DATE
    setting["A3"] = "判斷「逾期」用的日期。教材固定為 2026-04-15，實務上可改成 =TODAY()。"
    setting["A3"].font = NOTE_FONT
    set_widths(setting, [10, 14])
    wb.defined_names.add(DefinedName("基準日", attr_text="'設定'!$B$1"))

    cf = ws.conditional_formatting
    cf.add(f"I2:I{last}", CellIsRule(operator="lessThan", formula=["0"],
                                     font=Font(color="9C0006", bold=True),
                                     fill=PatternFill("solid", fgColor="FFC7CE")))
    cf.add(f"A2:J{last}", FormulaRule(formula=['AND($D2<基準日,$E2<>"已完成")'],
                                      fill=PatternFill("solid", fgColor="FFEB9C")))
    cf.add(f"E2:E{last}", FormulaRule(formula=['$E2="已完成"'], font=Font(color="006100"),
                                      fill=PatternFill("solid", fgColor="C6EFCE")))
    cf.add(f"F2:F{last}", DataBarRule(start_type="num", start_value=0, end_type="num",
                                      end_value=1, color="5B9BD5"))

    wb.save(lesson_path("07_條件式格式", "專案進度管控表.xlsx"))


# ---------------------------------------------------------------- 08
PRODUCTS = [
    ("P-001", "無線滑鼠", "電腦周邊", 590, 320),
    ("P-002", "機械鍵盤", "電腦周邊", 2490, 1450),
    ("P-003", "27 吋螢幕", "顯示器", 6990, 4800),
    ("P-004", "USB-C 集線器", "配件", 890, 420),
    ("P-005", "筆電支架", "配件", 1290, 610),
    ("P-006", "降噪耳機", "音訊", 3990, 2300),
    ("P-007", "網路攝影機", "視訊", 1790, 980),
    ("P-008", "藍牙喇叭", "音訊", 1590, 820),
]
ORDERS = [
    ("O-0001", date(2026, 1, 3), "P-003", 2), ("O-0002", date(2026, 1, 5), "P-001", 10),
    ("O-0003", date(2026, 1, 6), "P-006", 3), ("O-0004", date(2026, 1, 8), "P-002", 4),
    ("O-0005", date(2026, 1, 10), "P-004", 6), ("O-0006", date(2026, 1, 12), "P-999", 2),
    ("O-0007", date(2026, 1, 14), "P-005", 5), ("O-0008", date(2026, 1, 15), "P-007", 3),
    ("O-0009", date(2026, 1, 17), "P-008", 4), ("O-0010", date(2026, 1, 20), "P-003", 1),
    ("O-0011", date(2026, 1, 22), "P-001", 8), ("O-0012", date(2026, 1, 24), "P-006", 2),
    ("O-0013", date(2026, 1, 26), "P-002", 3), ("O-0014", date(2026, 1, 28), "P-004", 12),
    ("O-0015", date(2026, 1, 30), "P-008", 6),
]
NOT_FOUND = "查無此編號"


def build_08():
    from openpyxl.formatting.rule import FormulaRule

    wb = Workbook()
    orders = wb.active
    orders.title = "訂單明細"
    master = wb.create_sheet("產品主檔")
    summary = wb.create_sheet("訂單彙總")

    write_header(master, 1, ["產品編號", "產品名稱", "類別", "單價", "成本"])
    write_rows(master, 2, PRODUCTS, {4: AMOUNT, 5: AMOUNT})
    add_table(master, f"A1:E{1 + len(PRODUCTS)}", "T_Product")
    set_widths(master, [10, 14, 10, 9, 9])
    master.freeze_panes = "A2"

    write_header(orders, 1, ["訂單編號", "訂單日期", "產品編號", "數量",
                             "產品名稱", "類別", "單價", "小計"])

    def lookup(r, col, missing):
        # XLOOKUP 是 Excel 2021／365 的新函數，檔案內必須寫成 _xlfn.XLOOKUP
        return f"=_xlfn.XLOOKUP(C{r},T_Product[產品編號],T_Product[{col}],{missing})"

    rows = []
    for i, order in enumerate(ORDERS):
        r = 2 + i
        rows.append(order + (lookup(r, "產品名稱", f'"{NOT_FOUND}"'), lookup(r, "類別", '"—"'),
                             lookup(r, "單價", "0"), f"=D{r}*G{r}"))
    write_rows(orders, 2, rows, {2: DATE, 4: AMOUNT, 7: AMOUNT, 8: AMOUNT})
    last = 1 + len(rows)
    add_table(orders, f"A1:H{last}", "T_Order")
    set_widths(orders, [10, 12, 10, 7, 14, 10, 9, 10])
    orders.freeze_panes = "A2"
    orders.conditional_formatting.add(
        f"A2:H{last}", FormulaRule(formula=[f'$E2="{NOT_FOUND}"'], font=Font(color="9C0006"),
                                   fill=PatternFill("solid", fgColor="FFC7CE")))

    write_header(summary, 1, ["彙總項目", "數值", "公式說明"])
    items = [
        ("訂單筆數", "=COUNTA(T_Order[訂單編號])", "0", "數有幾筆訂單"),
        ("營收總計", "=SUM(T_Order[小計])", AMOUNT, "查無編號的訂單單價為 0，不計入營收"),
        ("查無編號筆數", f'=COUNTIF(T_Order[產品名稱],"{NOT_FOUND}")', "0", "不是 0 就要回頭檢查訂單"),
    ]
    for r, (name, formula, fmt, note) in enumerate(items, start=2):
        summary.cell(row=r, column=1, value=name).font = Font(name=FONT_NAME, bold=True)
        cell = summary.cell(row=r, column=2, value=formula)
        cell.number_format = fmt
        summary.cell(row=r, column=3, value=note).font = NOTE_FONT
    set_widths(summary, [14, 12, 36])

    wb.save(lesson_path("08_主檔與跨表查詢", "訂單與產品主檔.xlsx"))


# ---------------------------------------------------------------- 09, 12 共用
CHANNELS = ["門市", "官網", "電商平台"]
CATEGORIES = ["電腦周邊", "音訊", "配件"]
MONTHS = [date(2026, 1, 1), date(2026, 2, 1), date(2026, 3, 1)]


def make_sales(seed, n, with_cost=False):
    """產生固定（可重現）的第一季銷售明細。"""
    import random

    rng = random.Random(seed)
    rows = []
    for i in range(n):
        month = MONTHS[i * len(MONTHS) // n]
        d = date(month.year, month.month, rng.randint(1, 28))
        amount = rng.randint(50, 1500) * 10
        row = [f"S-{i + 1:04d}", d, rng.choice(CHANNELS), rng.choice(CATEGORIES), amount]
        if with_cost:
            row.append(round(amount * rng.uniform(0.52, 0.72) / 10) * 10)
        rows.append(tuple(row))
    rows.sort(key=lambda r: (r[1], r[0]))
    return [(f"S-{i + 1:04d}",) + r[1:] for i, r in enumerate(rows)]


SALES_09 = make_sales(seed=9, n=45)


def build_09():
    wb = Workbook()
    ws = wb.active
    ws.title = "銷售明細"
    write_header(ws, 1, ["訂單編號", "訂單日期", "通路", "產品類別", "銷售金額"])
    write_rows(ws, 2, SALES_09, {2: DATE, 5: AMOUNT})
    add_table(ws, f"A1:E{1 + len(SALES_09)}", "T_Sales")
    set_widths(ws, [10, 12, 10, 10, 11])
    ws.freeze_panes = "A2"

    s = wb.create_sheet("通路彙總")
    write_header(s, 1, ["通路", "訂單筆數", "銷售金額", "平均客單價"])
    for r, ch in enumerate(CHANNELS, start=2):
        s.cell(row=r, column=1, value=ch).font = Font(name=FONT_NAME, bold=True)
        s.cell(row=r, column=2, value=f"=COUNTIFS(T_Sales[通路],$A{r})")
        s.cell(row=r, column=3, value=f"=SUMIFS(T_Sales[銷售金額],T_Sales[通路],$A{r})").number_format = AMOUNT
        s.cell(row=r, column=4, value=f"=IF(B{r}=0,0,ROUND(C{r}/B{r},0))").number_format = AMOUNT
    total = 2 + len(CHANNELS)
    label(s, f"A{total}", "合計")
    s[f"B{total}"] = f"=SUM(B2:B{total - 1})"
    s[f"C{total}"] = f"=SUM(C2:C{total - 1})"
    s[f"C{total}"].number_format = AMOUNT
    s[f"D{total}"] = f"=IF(B{total}=0,0,ROUND(C{total}/B{total},0))"
    s[f"D{total}"].number_format = AMOUNT

    # 交叉矩陣：通路 × 月份（第 8 列為月份表頭，存成每月 1 日的日期）
    top = total + 3
    s.cell(row=top - 1, column=1, value="各通路每月銷售金額（SUMIFS 多條件加總）").font = Font(name=FONT_NAME, bold=True)
    write_header(s, top, ["通路＼月份"] + [""] * len(MONTHS) + ["合計"])
    for c, m in enumerate(MONTHS, start=2):
        cell = s.cell(row=top, column=c, value=m)
        cell.number_format = "yyyy-mm"
    for i, ch in enumerate(CHANNELS):
        r = top + 1 + i
        s.cell(row=r, column=1, value=ch).font = Font(name=FONT_NAME, bold=True)
        for c in range(2, 2 + len(MONTHS)):
            col = get_column_letter(c)
            s.cell(row=r, column=c, value=(
                f'=SUMIFS(T_Sales[銷售金額],T_Sales[通路],$A{r},'
                f'T_Sales[訂單日期],">="&{col}${top},T_Sales[訂單日期],"<"&EDATE({col}${top},1))'
            )).number_format = AMOUNT
        end = get_column_letter(1 + len(MONTHS))
        s.cell(row=r, column=2 + len(MONTHS), value=f"=SUM(B{r}:{end}{r})").number_format = AMOUNT
    set_widths(s, [14, 12, 12, 12, 12])
    wb.save(lesson_path("09_條件統計彙總", "多通路銷售彙總.xlsx"))


# ---------------------------------------------------------------- 10
RAW_MESSAGES = [
    "【研發部 林志遠】115/3/2 想買 27吋螢幕 2台，一台大約 NT$6,500",
    "行銷部張雅晴：3/3 需要 A4影印紙 10包 每包135元",
    "維運部-陳建宏 2026.03.05 網路線(5米) x6條 @450",
    "【管理部 劉美琪】115年3月8日 白板筆 4盒，每盒 160 元",
    "研發部 王思婷 2026/3/10 機械鍵盤 1 個 2,490",
    "業務部黃俊傑 3月12日 名片印製 3盒 單價$350",
    "【行銷部 張雅晴】115.03.15 展示架 2座 一座1,800元",
    "維運部 陳建宏：2026-03-18 UPS不斷電系統 1台 NT$4,200",
    "管理部-劉美琪 115/3/20 碳粉匣 3支，每支 1,150",
    "【業務部 黃俊傑】3/25 筆電支架 5 個 @1,290元",
]
CLEANED = [
    (1, date(2026, 3, 2), "研發部", "林志遠", "27 吋螢幕", 2, "台", 6500),
    (2, date(2026, 3, 3), "行銷部", "張雅晴", "A4 影印紙", 10, "包", 135),
    (3, date(2026, 3, 5), "維運部", "陳建宏", "網路線（5 米）", 6, "條", 450),
    (4, date(2026, 3, 8), "管理部", "劉美琪", "白板筆", 4, "盒", 160),
    (5, date(2026, 3, 10), "研發部", "王思婷", "機械鍵盤", 1, "個", 2490),
    (6, date(2026, 3, 12), "業務部", "黃俊傑", "名片印製", 3, "盒", 350),
    (7, date(2026, 3, 15), "行銷部", "張雅晴", "展示架", 2, "座", 1800),
    (8, date(2026, 3, 18), "維運部", "陳建宏", "UPS 不斷電系統", 1, "台", 4200),
    (9, date(2026, 3, 20), "管理部", "劉美琪", "碳粉匣", 3, "支", 1150),
    (10, date(2026, 3, 25), "業務部", "黃俊傑", "筆電支架", 5, "個", 1290),
]
CLEANING_RULES = [
    ("申請日期", "115/3/2、2026.03.05、115年3月8日、3/3", "民國年 + 1911 轉西元；沒寫年份的視為 2026 年；統一存成日期 yyyy-mm-dd"),
    ("部門、申請人", "【研發部 林志遠】、行銷部張雅晴：、維運部-陳建宏", "拆成兩欄；去掉【】、：、- 等符號與多餘空白"),
    ("品項", "27吋螢幕、A4影印紙", "保留品名；中英文與數字之間補空格，括號統一為全形"),
    ("數量、單位", "2台、x6條、1 個", "數字放「數量」（數值），單位放「單位」（文字），一格一事"),
    ("單價", "NT$6,500、135元、@450、$350", "去掉 NT$、$、元、@、逗號，存成數值，顯示格式 #,##0"),
    ("小計", "原始訊息沒有", "清洗後用公式 =數量×單價 計算，不由 AI 自己算好填入"),
]


def build_10():
    wb = Workbook()
    raw = wb.active
    raw.title = "原始訊息"
    write_header(raw, 1, ["原始編號", "原始訊息（LINE 群組複製）"])
    write_rows(raw, 2, list(enumerate(RAW_MESSAGES, start=1)))
    set_widths(raw, [10, 62])
    raw.freeze_panes = "A2"

    clean = wb.create_sheet("清洗後")
    write_header(clean, 1, ["原始編號", "申請日期", "部門", "申請人", "品項", "數量", "單位", "單價", "小計"])
    rows = [row + (f"=F{2 + i}*H{2 + i}",) for i, row in enumerate(CLEANED)]
    write_rows(clean, 2, rows, {2: DATE, 6: AMOUNT, 8: AMOUNT, 9: AMOUNT})
    add_table(clean, f"A1:I{1 + len(rows)}", "T_Purchase")
    set_widths(clean, [10, 12, 9, 9, 18, 7, 7, 9, 10])
    clean.freeze_panes = "A2"

    log = wb.create_sheet("清洗規則")
    write_header(log, 1, ["欄位", "原始樣貌", "清洗規則"])
    write_rows(log, 2, CLEANING_RULES)
    for r in range(2, 2 + len(CLEANING_RULES)):
        for c in (2, 3):
            log.cell(row=r, column=c).alignment = Alignment(wrap_text=True, vertical="top")
    set_widths(log, [14, 36, 52])

    wb.save(lesson_path("10_髒資料清洗", "採購需求清洗.xlsx"))


# ---------------------------------------------------------------- 11
MONTHLY_REVENUE = [
    (820000, 310000), (760000, 290000), (880000, 350000), (840000, 380000),
    (900000, 420000), (870000, 460000), (910000, 500000), (890000, 530000),
    (860000, 560000), (930000, 600000), (1050000, 720000), (1120000, 810000),
]


def _axes(chart, x_title, y_title):
    chart.x_axis.title = x_title
    chart.y_axis.title = y_title
    chart.y_axis.scaling.min = 0  # 數值軸從 0 開始（第二部選圖規則）
    chart.y_axis.number_format = "#,##0"
    chart.x_axis.delete = False
    chart.y_axis.delete = False
    chart.width, chart.height = 18, 8


def build_11():
    from openpyxl.chart import BarChart, LineChart, PieChart, Reference
    from openpyxl.chart.label import DataLabelList

    wb = Workbook()
    ws = wb.active
    ws.title = "月營收"
    write_header(ws, 1, ["月份", "門市營收", "線上營收", "合計營收"])
    rows = [(date(2025, m, 1), store, online, f"=B{m + 1}+C{m + 1}")
            for m, (store, online) in enumerate(MONTHLY_REVENUE, start=1)]
    write_rows(ws, 2, rows, {1: "yyyy-mm", 2: AMOUNT, 3: AMOUNT, 4: AMOUNT})
    set_widths(ws, [10, 12, 12, 12])
    ws.freeze_panes = "A2"

    q = wb.create_sheet("季彙總")
    write_header(q, 1, ["季別", "門市營收", "線上營收", "合計營收"])
    for i in range(4):
        r, a, b = 2 + i, 2 + i * 3, 4 + i * 3
        q.cell(row=r, column=1, value=f"Q{i + 1}").font = Font(name=FONT_NAME, bold=True)
        for c, col in ((2, "B"), (3, "C")):
            q.cell(row=r, column=c, value=f"=SUM('月營收'!{col}{a}:{col}{b})").number_format = AMOUNT
        q.cell(row=r, column=4, value=f"=B{r}+C{r}").number_format = AMOUNT
    write_header(q, 8, ["通路", "全年營收"])
    for r, (name, col) in ((9, ("門市", "B")), (10, ("線上", "C"))):
        q.cell(row=r, column=1, value=name).font = Font(name=FONT_NAME, bold=True)
        q.cell(row=r, column=2, value=f"=SUM({col}2:{col}5)").number_format = AMOUNT
    set_widths(q, [10, 12, 12, 12])

    charts = wb.create_sheet("圖表")
    charts["A1"] = "三種問題，三種圖：趨勢 → 折線圖、比較 → 長條圖、組成 → 圓餅圖"
    charts["A1"].font = Font(name=FONT_NAME, bold=True, size=13)

    line = LineChart()
    line.title = "2025 每月合計營收（趨勢）"
    line.add_data(Reference(ws, min_col=4, min_row=1, max_row=13), titles_from_data=True)
    line.set_categories(Reference(ws, min_col=1, min_row=2, max_row=13))
    line.x_axis.number_format = "yyyy-mm"
    line.legend = None
    _axes(line, "月份", "營收（元）")
    charts.add_chart(line, "A3")

    bar = BarChart()
    bar.type = "col"
    bar.title = "各季門市與線上營收（比較）"
    bar.add_data(Reference(q, min_col=2, max_col=3, min_row=1, max_row=5), titles_from_data=True)
    bar.set_categories(Reference(q, min_col=1, min_row=2, max_row=5))
    _axes(bar, "季別", "營收（元）")
    charts.add_chart(bar, "A20")

    pie = PieChart()
    pie.title = "全年營收：門市 vs. 線上（組成）"
    pie.add_data(Reference(q, min_col=2, min_row=8, max_row=10), titles_from_data=True)
    pie.set_categories(Reference(q, min_col=1, min_row=9, max_row=10))
    pie.dataLabels = DataLabelList()
    pie.dataLabels.showPercent = True
    pie.width, pie.height = 12, 8
    charts.add_chart(pie, "A37")

    wb.save(lesson_path("11_圖表", "月營收圖表.xlsx"))


# ---------------------------------------------------------------- 12
SALES_12 = make_sales(seed=12, n=60, with_cost=True)


def dashboard_numbers():
    """儀表板的關鍵數字，供重點文字與驗收腳本共用。"""
    revenue = sum(r[4] for r in SALES_12)
    profit = sum(r[4] - r[5] for r in SALES_12)
    by_month = {m: sum(r[4] for r in SALES_12 if r[1].month == m.month) for m in MONTHS}
    by_channel = {ch: (sum(r[4] for r in SALES_12 if r[2] == ch),
                       sum(r[4] - r[5] for r in SALES_12 if r[2] == ch)) for ch in CHANNELS}
    return {"revenue": revenue, "profit": profit, "orders": len(SALES_12),
            "by_month": by_month, "by_channel": by_channel}


def build_12():
    from openpyxl.chart import BarChart, Reference

    wb = Workbook()
    dash = wb.active
    dash.title = "儀表板"
    calc = wb.create_sheet("計算")
    data = wb.create_sheet("銷售明細")

    # 資料層
    write_header(data, 1, ["訂單編號", "訂單日期", "通路", "產品類別", "銷售金額", "成本", "毛利"])
    rows = [r + (f"=E{2 + i}-F{2 + i}",) for i, r in enumerate(SALES_12)]
    write_rows(data, 2, rows, {2: DATE, 5: AMOUNT, 6: AMOUNT, 7: AMOUNT})
    add_table(data, f"A1:G{1 + len(rows)}", "T_Sales12")
    set_widths(data, [10, 12, 10, 10, 11, 10, 10])
    data.freeze_panes = "A2"

    # 展示層的計算區：每月與每通路
    write_header(calc, 1, ["月份", "營收", "毛利"])
    for i, m in enumerate(MONTHS):
        r = 2 + i
        calc.cell(row=r, column=1, value=m).number_format = "yyyy-mm"
        cond = f'T_Sales12[訂單日期],">="&$A{r},T_Sales12[訂單日期],"<"&EDATE($A{r},1)'
        calc.cell(row=r, column=2, value=f"=SUMIFS(T_Sales12[銷售金額],{cond})").number_format = AMOUNT
        calc.cell(row=r, column=3, value=f"=SUMIFS(T_Sales12[毛利],{cond})").number_format = AMOUNT
    write_header(calc, 7, ["通路", "營收", "毛利", "毛利率"])
    for i, ch in enumerate(CHANNELS):
        r = 8 + i
        calc.cell(row=r, column=1, value=ch)
        calc.cell(row=r, column=2, value=f"=SUMIFS(T_Sales12[銷售金額],T_Sales12[通路],$A{r})").number_format = AMOUNT
        calc.cell(row=r, column=3, value=f"=SUMIFS(T_Sales12[毛利],T_Sales12[通路],$A{r})").number_format = AMOUNT
        calc.cell(row=r, column=4, value=f"=IF(B{r}=0,0,C{r}/B{r})").number_format = "0.0%"
    set_widths(calc, [10, 12, 12, 10])

    # 展示層：儀表板（只有這裡允許合併儲存格）
    dash.merge_cells("A1:L1")
    dash["A1"] = "2026 第一季全通路營運儀表板"
    dash["A1"].font = Font(name=FONT_NAME, bold=True, size=18, color="FFFFFF")
    dash["A1"].fill = HEADER_FILL
    dash["A1"].alignment = Alignment(horizontal="center", vertical="center")
    dash.row_dimensions[1].height = 34

    cards = [
        ("總營收", "=SUM(T_Sales12[銷售金額])", AMOUNT),
        ("總毛利", "=SUM(T_Sales12[毛利])", AMOUNT),
        ("毛利率", "=IF(SUM(T_Sales12[銷售金額])=0,0,SUM(T_Sales12[毛利])/SUM(T_Sales12[銷售金額]))", "0.0%"),
        ("訂單筆數", "=COUNTA(T_Sales12[訂單編號])", "#,##0"),
    ]
    card_fill = PatternFill("solid", fgColor="DDEBF7")
    for i, (name, formula, fmt) in enumerate(cards):
        c1 = 1 + i * 3
        left, right = get_column_letter(c1), get_column_letter(c1 + 2)
        dash.merge_cells(f"{left}3:{right}3")
        dash.merge_cells(f"{left}4:{right}5")
        title = dash[f"{left}3"]
        title.value = name
        title.font = Font(name=FONT_NAME, bold=True, color="1F4E79")
        value = dash[f"{left}4"]
        value.value = formula
        value.number_format = fmt
        value.font = Font(name=FONT_NAME, bold=True, size=20, color="1F4E79")
        for ref in (title, value):
            ref.fill = card_fill
            ref.alignment = Alignment(horizontal="center", vertical="center")
    for col in range(1, 13):
        dash.column_dimensions[get_column_letter(col)].width = 9

    month_chart = BarChart()
    month_chart.type = "col"
    month_chart.title = "每月營收"
    month_chart.add_data(Reference(calc, min_col=2, min_row=1, max_row=4), titles_from_data=True)
    month_chart.set_categories(Reference(calc, min_col=1, min_row=2, max_row=4))
    month_chart.x_axis.number_format = "yyyy-mm"
    month_chart.legend = None
    _axes(month_chart, "月份", "營收（元）")
    month_chart.width, month_chart.height = 13, 7.5
    dash.add_chart(month_chart, "A7")

    channel_chart = BarChart()
    channel_chart.type = "bar"
    channel_chart.title = "各通路營收與毛利"
    channel_chart.add_data(Reference(calc, min_col=2, max_col=3, min_row=7, max_row=10), titles_from_data=True)
    channel_chart.set_categories(Reference(calc, min_col=1, min_row=8, max_row=10))
    _axes(channel_chart, "通路", "金額（元）")
    channel_chart.width, channel_chart.height = 13, 7.5
    dash.add_chart(channel_chart, "G7")

    nums = dashboard_numbers()
    best_month = max(nums["by_month"], key=nums["by_month"].get)
    best_channel = max(CHANNELS, key=lambda ch: nums["by_channel"][ch][0])
    rates = {ch: p / r for ch, (r, p) in nums["by_channel"].items()}
    best_rate = max(rates, key=rates.get)
    notes = [
        "📝 本季重點（AI 依示範資料撰寫，數字可對照上方卡片與「計算」工作表）",
        f"1. 營收最高的月份是 {best_month.month} 月，共 {nums['by_month'][best_month]:,} 元。",
        f"2. 營收最高的通路是「{best_channel}」，共 {nums['by_channel'][best_channel][0]:,} 元。",
        f"3. 毛利率最高的通路是「{best_rate}」（{rates[best_rate]:.1%}），可優先檢視其定價做法。",
        "⚠️ 限制：只有一季資料，不能推論長期趨勢；毛利未扣除運費與人事成本。",
    ]
    for i, text in enumerate(notes):
        cell = dash.cell(row=23 + i, column=1, value=text)
        cell.font = Font(name=FONT_NAME, bold=(i == 0))

    wb.save(lesson_path("12_主管儀表板總整", "全通路營運儀表板.xlsx"))


BUILDERS = [build_01, build_02, build_03, build_04, build_05, build_06,
            build_07, build_08, build_09, build_10, build_11, build_12]

if __name__ == "__main__":
    for build in BUILDERS:
        build()
        print(f"✅ {build.__name__}")
