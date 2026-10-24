"""自動驗收第一部各課 xlsx 是否符合規範書。

檢查：無合併儲存格、日期欄是真日期、數值欄是數字、公式欄是公式、
表格／資料驗證／條件格式／圖表數量、新函數有 _xlfn. 前綴、關鍵數值以 Python 驗算。
執行：uv run python scripts/check_part1_workbooks.py
"""

import re
import sys
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent / "用AI生成Excel"

# 每課的期望值：sheet -> {"rows": (起, 迄), "date"/"number"/"text"/"formula": 欄, 各類數量}
SPECS = {
    "01_儲存格與資料型態/社團收支紀錄.xlsx": {
        "收支紀錄": {"rows": (2, 11), "date": "A", "number": "D", "text": "F"},
        "型態對照": {"rows": (2, 2), "number": "C", "text": "B"},
    },
    "02_公式與參照/商品報價單.xlsx": {
        "報價單": {"rows": (7, 14), "number": "CD", "formula": "EFG", "absolute": ("F", "$B$1")},
        "錯誤示範_相對參照": {"rows": (7, 14), "number": "CD", "formula": "EFG"},
    },
    "03_常用函數/班級成績表.xlsx": {
        "成績": {"rows": (2, 21), "formula": "FGH", "text": "A"},
        "各科統計": {"rows": (2, 6), "formula": "BCD"},
    },
    "04_一列一筆與結構化表格/部門費用報銷表.xlsx": {
        "報銷明細": {"rows": (2, 11), "date": "B", "number": "GH", "formula": "I", "tables": 1},
        "報銷彙總": {"rows": (2, 4), "formula": "B"},
    },
    "05_邏輯判斷/業績獎金判定.xlsx": {
        "業績": {"rows": (2, 13), "number": "DE", "formula": "GHIJ", "tables": 1},
        "獎金規則": {"rows": (2, 4), "number": "B"},
    },
    "06_資料驗證/請假申請表.xlsx": {
        "請假紀錄": {"rows": (2, 11), "date": "BFG", "number": "H", "tables": 1, "validations": 6},
        "選項清單": {"rows": (2, 4), "text": "ABC"},
    },
    "07_條件式格式/專案進度管控表.xlsx": {
        "專案進度": {"rows": (2, 13), "date": "D", "number": "FGH", "formula": "IJ",
                 "tables": 1, "cf_rules": 4},
        "設定": {"rows": (1, 1), "date": "B"},
    },
    "08_主檔與跨表查詢/訂單與產品主檔.xlsx": {
        "訂單明細": {"rows": (2, 16), "date": "B", "number": "D", "formula": "EFGH",
                 "tables": 1, "cf_rules": 1},
        "產品主檔": {"rows": (2, 9), "number": "DE", "tables": 1},
        "訂單彙總": {"rows": (2, 4), "formula": "B"},
    },
    "09_條件統計彙總/多通路銷售彙總.xlsx": {
        "銷售明細": {"rows": (2, 46), "date": "B", "number": "E", "tables": 1},
        "通路彙總": {"rows": (2, 5), "formula": "BCD"},
    },
    "10_髒資料清洗/採購需求清洗.xlsx": {
        "原始訊息": {"rows": (2, 11), "text": "B"},
        "清洗後": {"rows": (2, 11), "date": "B", "number": "FH", "text": "G", "formula": "I", "tables": 1},
    },
    "11_圖表/月營收圖表.xlsx": {
        "月營收": {"rows": (2, 13), "date": "A", "number": "BC", "formula": "D"},
        "季彙總": {"rows": (2, 5), "formula": "BCD"},
        "圖表": {"rows": (1, 1), "charts": 3},
    },
    "12_主管儀表板總整/全通路營運儀表板.xlsx": {
        "銷售明細": {"rows": (2, 61), "date": "B", "number": "EF", "formula": "G", "tables": 1},
        "計算": {"rows": (2, 4), "date": "A", "formula": "BC"},
        "儀表板": {"rows": (4, 4), "formula": "ADGJ", "charts": 2, "allow_merge": True},
    },
}

# Excel 2010 之後的新函數，檔案內必須加 _xlfn. 前綴，否則開檔會顯示 #NAME?
NEW_FUNCTIONS = re.compile(r"(?<![A-Z_.])(XLOOKUP|IFS|MAXIFS|MINIFS|SWITCH|TEXTJOIN|FILTER|UNIQUE)\(")


def count(ws, spec, key, actual, errors, label):
    if key in spec and actual != spec[key]:
        errors.append(f"{ws.title}：{label} {actual} 個，應為 {spec[key]} 個")


def check_file(rel, sheets):
    errors = []
    wb = load_workbook(ROOT / rel)
    for name, spec in sheets.items():
        if name not in wb.sheetnames:
            errors.append(f"缺少工作表「{name}」")
            continue
        ws = wb[name]
        if ws.merged_cells.ranges and not spec.get("allow_merge"):
            errors.append(f"{name}：有合併儲存格 {ws.merged_cells.ranges}")
        count(ws, spec, "tables", len(ws.tables), errors, "表格")
        count(ws, spec, "validations", len(ws.data_validations.dataValidation), errors, "資料驗證")
        count(ws, spec, "cf_rules", sum(len(cf.rules) for cf in ws.conditional_formatting), errors, "條件式格式規則")
        count(ws, spec, "charts", len(ws._charts), errors, "圖表")

        start, end = spec["rows"]
        for r in range(start, end + 1):
            for col in spec.get("date", ""):
                v = ws[f"{col}{r}"].value
                if not isinstance(v, (date, datetime)):
                    errors.append(f"{name}!{col}{r} 應為日期，實際為 {type(v).__name__}：{v!r}")
            for col in spec.get("number", ""):
                v = ws[f"{col}{r}"].value
                if not isinstance(v, (int, float)) or isinstance(v, bool):
                    errors.append(f"{name}!{col}{r} 應為數值，實際為 {v!r}")
            for col in spec.get("text", ""):
                v = ws[f"{col}{r}"].value
                if not isinstance(v, str):
                    errors.append(f"{name}!{col}{r} 應為文字，實際為 {v!r}")
            for col in spec.get("formula", ""):
                v = ws[f"{col}{r}"].value
                if not (isinstance(v, str) and v.startswith("=")):
                    errors.append(f"{name}!{col}{r} 應為公式，實際為 {v!r}")
            if "absolute" in spec:
                col, ref = spec["absolute"]
                if ref not in str(ws[f"{col}{r}"].value):
                    errors.append(f"{name}!{col}{r} 應使用絕對參照 {ref}")

        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    m = NEW_FUNCTIONS.search(cell.value)
                    if m:
                        errors.append(f"{name}!{cell.coordinate} 的 {m.group(1)} 缺少 _xlfn. 前綴")

        for table in ws.tables.values():
            # 總計等彙總值不可落在表格範圍內（避免排序、篩選時被當成一筆資料）
            last_row = int(re.sub(r"\D", "", table.ref.split(":")[1]))
            if last_row != end:
                errors.append(f"{name}：表格 {table.displayName} 範圍 {table.ref} 與資料列不符")
    return errors


def recompute():
    """不靠 Excel 重算，用 Python 驗算關鍵數字，供規範書 C7 引用。"""
    sys.path.insert(0, str(Path(__file__).resolve().parent))
    import build_part1_workbooks as b

    subtotal = sum(p * q for _, _, p, q in b.QUOTE_ITEMS)
    tax = sum(int(p * q * 0.05 + 0.5) for _, _, p, q in b.QUOTE_ITEMS)

    def grade(rate):
        return next((g for g, low, _ in b.BONUS_RULES if rate >= low), "無")

    grades = [grade(actual / target) for _, _, _, target, actual, _ in b.SALES_REPS]
    extra = sum(1 for (_, _, _, t, a, full) in b.SALES_REPS if a / t >= 1 and full == "是")
    overdue = [t[0] for t in b.TASKS if t[3] < b.BASE_DATE and t[4] != "已完成"]
    over_budget = [t[0] for t in b.TASKS if t[7] > t[6]]
    prices = {p[0]: p[3] for p in b.PRODUCTS}
    nums = b.dashboard_numbers()
    return {
        "02 未稅合計／稅額合計／含稅總額": f"{subtotal:,}／{tax:,}／{subtotal + tax:,}",
        "04 報銷總計金額": f"{sum(r[6] * r[7] for r in b.EXPENSES):,}",
        "05 獎金等級 A/B/C/無": "/".join(str(grades.count(g)) for g in ("A", "B", "C", "無")),
        "05 全勤加碼人數": extra,
        "07 逾期任務／超支任務": f"{overdue}／{over_budget}",
        "08 營收總計": f"{sum(q * prices.get(c, 0) for _, _, c, q in b.ORDERS):,}",
        "09 各通路筆數與金額": {ch: (sum(1 for r in b.SALES_09 if r[2] == ch),
                              f"{sum(r[4] for r in b.SALES_09 if r[2] == ch):,}") for ch in b.CHANNELS},
        "10 小計合計": f"{sum(r[5] * r[7] for r in b.CLEANED):,}",
        "11 門市／線上全年": f"{sum(s for s, _ in b.MONTHLY_REVENUE):,}／{sum(o for _, o in b.MONTHLY_REVENUE):,}",
        "12 總營收／總毛利／毛利率／筆數": f"{nums['revenue']:,}／{nums['profit']:,}／"
                                 f"{nums['profit'] / nums['revenue']:.1%}／{nums['orders']}",
    }


def main():
    failed = False
    for rel, sheets in SPECS.items():
        errors = check_file(rel, sheets)
        if errors:
            failed = True
            print(f"❌ {rel}")
            for e in errors:
                print(f"   - {e}")
        else:
            print(f"✅ {rel}")
    print("\n📐 Python 驗算：")
    for k, v in recompute().items():
        print(f"   {k}：{v}")
    sys.exit(1 if failed else 0)


if __name__ == "__main__":
    main()
